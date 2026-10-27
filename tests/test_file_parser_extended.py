"""
Unit Tests for Ticket 1.8: Extended File Parser & Structure Extraction.
Tests PyMuPDF/fitz extraction, DOCX headings and tables, progress callbacks, and boundary detection.
"""

import io
import pytest
from src.web.file_parser import (
    extract_text_from_file,
    detect_structure_boundaries,
    sanitize_filename,
    format_file_size,
    validate_file,
)


def test_pdf_extraction_with_pymupdf():
    """Verify PyMuPDF extracts text and detects headings in PDF bytes."""
    import fitz

    # Create a small in-memory PDF with headings and content
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "CHƯƠNG 1. TỔNG QUAN KIẾN TRÚC SA-CMS", fontsize=16)
    page.insert_text((50, 90), "Hệ thống SA-CMS phân cấp bộ nhớ thành 3 thang thời gian tương ứng với cấu trúc tài liệu.", fontsize=11)
    page.insert_text((50, 130), "MỤC 1.1. CƠ CHẾ GATING ĐỘNG", fontsize=14)
    page.insert_text((50, 160), "Cổng hòa trộn g_l(x) điều tiết dòng thông tin từ bộ nhớ tham số vào biểu diễn ẩn.", fontsize=11)

    pdf_bytes = doc.tobytes()
    doc.close()

    text, meta = extract_text_from_file("sacms_overview.pdf", pdf_bytes)
    assert "CHƯƠNG 1" in text or "SA-CMS" in text
    assert meta["extension"] == ".pdf"
    assert meta["word_count"] > 10
    assert "structure_boundaries" in meta


def test_docx_structure_and_table_extraction():
    """Verify python-docx parser extracts heading styles and table rows."""
    import docx

    doc = docx.Document()
    doc.add_heading("Báo cáo Khoa học NCKH", level=1)
    doc.add_paragraph("Nội dung giới thiệu bài báo nghiên cứu.")
    doc.add_heading("Kết quả Thực nghiệm", level=2)
    doc.add_paragraph("Bảng số liệu kiểm chuẩn Token F1:")

    # Add table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Phương pháp"
    table.cell(0, 1).text = "Token F1"
    table.cell(1, 0).text = "P1 (SA-CMS)"
    table.cell(1, 1).text = "0.0421"

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    text, meta = extract_text_from_file("report.docx", docx_bytes)
    assert "# Báo cáo Khoa học NCKH" in text
    assert "## Kết quả Thực nghiệm" in text
    assert "| Phương pháp | Token F1 |" in text
    assert "| P1 (SA-CMS) | 0.0421 |" in text
    assert meta["extension"] == ".docx"


def test_progress_callback():
    """Verify on_progress callback receives updates from 0 to 100%."""
    progress_records = []

    def progress_tracker(pct: float, msg: str):
        progress_records.append((pct, msg))

    sample_text = b"# Chapter 1\n\nParagraph 1 text content.\n\n# Chapter 2\n\nParagraph 2 text content."
    text, meta = extract_text_from_file("notes.txt", sample_text, on_progress=progress_tracker)

    assert len(progress_records) >= 2
    assert progress_records[-1][0] == 100.0
    assert "Hoàn thành" in progress_records[-1][1]


def test_detect_structure_boundaries():
    """Verify boundary detection correctly counts Level 1, 2, and 3 divisions."""
    doc_text = (
        "# Phần 1: Giới thiệu\n\n"
        "Đoạn văn đầu tiên trình bày bối cảnh nghiên cứu.\n\n"
        "Đoạn văn thứ hai mô tả các thách thức của mô hình truyền thống.\n\n"
        "## Mục 1.1: Khảo sát lý thuyết\n\n"
        "Đoạn văn thứ ba phân tích cơ chế Attention và Nested Learning.\n\n"
        "# Phần 2: Phương pháp Đề xuất\n\n"
        "Đoạn văn thứ tư trình bày chi tiết công thức cập nhật trọng số.\n"
    )

    boundaries = detect_structure_boundaries(doc_text)
    assert boundaries["level_1_paragraphs"] >= 4
    assert boundaries["level_2_sections"] >= 3
    assert boundaries["level_3_documents"] == 1
    assert boundaries["boundary_indicators"]["b3"] == 1
